#!/usr/bin/env python3
"""UNOFFICIAL UniFi OS Server local-admin migration proof of concept.

No direct database writes. No official Ubiquiti API contract is claimed.
Tested workflow: UniFi OS Server 5.1.42, Network 10.5.67 (October 2026).
Use only with authorization and verified backups; inspect planned rights carefully.
"""
import argparse
import base64
import binascii
import csv
import hashlib
import json
import os
import re
import time
from collections import Counter, defaultdict
from pathlib import Path
from urllib.parse import urlparse

ROLE_PATH = "/proxy/users/api/v2/custom_role"
USER_PATH = "/proxy/users/api/v2/user"
READ_PATH = "/proxy/users/api/v2/custom_roles"
HEX24 = re.compile(r"^[a-f0-9]{24}$")
UUID = re.compile(r"^[a-f0-9]{8}-(?:[a-f0-9]{4}-){3}[a-f0-9]{12}$", re.I)
USER_RE = re.compile(r"^[A-Za-z0-9._-]{1,64}$")


def canon(s):
    return str(s or "").strip().casefold()


def hkey(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").casefold())


def sheet_rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    ws.reset_dimensions()
    it = ws.iter_rows(values_only=True)
    try:
        headers = next(it)
    except StopIteration:
        raise ValueError("Empty workbook")
    out = [dict(zip(headers, row)) for row in it]
    wb.close()
    return out


def read_people(path):
    if path.suffix.lower() != ".xlsx":
        raise ValueError("People input must be a .xlsx workbook")
    rows = sheet_rows(path)
    def field(row, *names):
        fields = {hkey(k): v for k, v in row.items() if k is not None}
        for name in names:
            if hkey(name) in fields:
                return fields[hkey(name)]
        return None
    result = {}
    for lineno, row in enumerate(rows, 2):
        username = str(field(row, "username", "Uporabniško ime", "Uporabnisko ime") or "").strip()
        if not username:
            continue
        key = canon(username)
        if key in result:
            raise ValueError(f"Duplicate username at Excel row {lineno}: {username}")
        def s(*names):
            value = field(row, *names)
            return str(value).strip() if value is not None else ""
        result[key] = {
            "username": username,
            "first_name": s("first_name", "Ime"),
            "last_name": s("last_name", "Priimek"),
            "email": s("email", "E-pošta", "E-posta"),
            "legacy_username": s("legacy_username", "Staro uporabniško ime", "Staro uporabnisko ime"),
            "site_count": s("site_count", "Število mest", "Stevilo mest"),
            "mode": s("mode", "Način", "Nacin").casefold() or "import",
            "password": field(row, "password", "Geslo"),
        }
    if not result:
        raise ValueError("No users found in workbook")
    return result


def read_legacy(path):
    result = defaultdict(lambda: {"email": "", "sites": {}})
    counts = Counter()
    summary = None
    with path.open(encoding="utf-8") as f:
        for lineno, line in enumerate(f, 1):
            if not line.strip():
                continue
            data = json.loads(line)
            if data.get("_export_summary") is True:
                if summary is not None:
                    raise ValueError("Duplicate legacy export summary")
                summary = data
                continue
            if summary is not None:
                raise ValueError("Legacy export has records after summary")
            name = str(data.get("admin_name") or "").strip()
            sid = str(data.get("site_id") or "").lower()
            aid = str(data.get("admin_id") or "").lower()
            role = str(data.get("role") or "").lower()
            if not name or not HEX24.fullmatch(sid) or not HEX24.fullmatch(aid):
                raise ValueError(f"Invalid admin/site ID at legacy line {lineno}")
            if role not in {"admin", "readonly", "nobody"}:
                raise ValueError(f"Unknown role {role!r} at legacy line {lineno}")
            entry = result[canon(name)]
            if entry.get("admin_id") and entry["admin_id"] != aid:
                raise ValueError(f"Username linked to two legacy IDs: {name}")
            entry["admin_id"] = aid
            entry["username"] = name
            email = str(data.get("admin_email") or "").strip()
            if entry["email"] and email and entry["email"] != email:
                raise ValueError(f"Conflicting email for {name}")
            if email:
                entry["email"] = email
            if sid in entry["sites"]:
                raise ValueError(f"Duplicate legacy privilege for {name}/{sid}")
            entry["sites"][sid] = {
                "role": role,
                "name": str(data.get("site_name") or ""),
                "desc": str(data.get("site_description") or ""),
            }
            counts[role] += 1
    if not result or summary is None:
        raise ValueError("Incomplete legacy export: summary line is missing")
    if summary.get("privilege_count") != sum(counts.values()):
        raise ValueError("Legacy privilege count does not match exporter summary")
    if summary.get("admin_count") != len(result):
        raise ValueError("Legacy admin count mismatch (some admins may lack privileges)")
    return result, counts


def read_target_sites(path):
    result = {}
    summary = None
    with path.open(encoding="utf-8") as f:
        for lineno, line in enumerate(f, 1):
            if not line.strip():
                continue
            data = json.loads(line)
            if data.get("_export_summary") is True:
                if summary is not None:
                    raise ValueError("Duplicate target export summary")
                summary = data
                continue
            if summary is not None:
                raise ValueError("Target export has records after summary")
            sid = str(data.get("site_id") or "").lower()
            if not HEX24.fullmatch(sid) or sid in result:
                raise ValueError(f"Invalid/duplicate target site ID at line {lineno}")
            result[sid] = {
                "name": str(data.get("site_name") or ""),
                "desc": str(data.get("site_description") or ""),
            }
    if not result or summary is None:
        raise ValueError("Incomplete target sites export: summary line is missing")
    if summary.get("site_count") != len(result):
        raise ValueError("Target site count does not match exporter summary")
    return result


def prepare(people, legacy, target, manual_usernames):
    rows = []
    old_mappings = defaultdict(list)
    for key in sorted(people):
        p = people[key]
        old_key = canon(p["legacy_username"]) or key
        old = legacy.get(old_key)
        reasons, warnings, formulas, sites = [], [], [], []
        if p["mode"] not in {"import", "manual"}:
            reasons.append("mode must be import or manual")
        if key == "admin":
            reasons.append("protected target Owner username: admin")
        if key in manual_usernames and p["mode"] != "manual":
            reasons.append("configured Super Admin must use mode=manual")
        if old is None:
            reasons.append("No matching legacy username; explicitly set legacy_username")
        else:
            old_mappings[old_key].append(key)
            if old_key != key:
                warnings.append(f"Renamed from legacy account: {old['username']}")
            if old["email"] and p["email"] and canon(old["email"]) != canon(p["email"]):
                warnings.append("Email differs from legacy; review identity mapping")
            for sid, spec in sorted(old["sites"].items()):
                if canon(spec["name"]) == "super" or spec["role"] == "nobody":
                    continue
                new_site = target.get(sid)
                if new_site is None:
                    reasons.append("Missing target site ID: " + sid)
                    continue
                if spec["name"] and new_site["name"] and spec["name"] != new_site["name"]:
                    reasons.append("Site ID matches but internal name differs: " + sid)
                if spec["desc"] and new_site["desc"] and spec["desc"] != new_site["desc"]:
                    warnings.append("Site description changed: " + sid)
                formulas.append(f"s-{sid}-{spec['role']}")
                sites.append((spec["desc"] or spec["name"] or sid, spec["role"]))
        if p["mode"] == "import":
            if not USER_RE.fullmatch(p["username"]):
                reasons.append("Invalid username characters/length")
            if not p["first_name"] or not p["last_name"]:
                reasons.append("First/last name missing")
            if p["email"] and not re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", p["email"]):
                reasons.append("Invalid email")
            if not isinstance(p["password"], str) or len(p["password"]) < 12:
                reasons.append("Password missing or shorter than 12 characters")
            elif p["password"] != p["password"].strip():
                reasons.append("Password has leading/trailing whitespace")
            if not formulas:
                reasons.append("No positive site privileges")
            if old and p["site_count"]:
                try:
                    if int(p["site_count"]) != len(formulas):
                        reasons.append("Declared site count does not match legacy privileges")
                except ValueError:
                    reasons.append("Invalid site count")
        rows.append({"key": key, "people": p, "legacy_key": old_key,
                     "legacy_username": old.get("username") if old else "?",
                     "formulas": formulas, "sites": sites,
                     "reasons": reasons, "warnings": warnings})
    for row in rows:
        if len(old_mappings[row["legacy_key"]]) > 1:
            row["reasons"].append("Multiple new accounts mapped to one legacy account")
    return rows


def fingerprint(row):
    obj = {"username": row["people"]["username"],
           "legacy_username": row["legacy_username"],
           "first_name": row["people"]["first_name"],
           "last_name": row["people"]["last_name"],
           "email": row["people"]["email"],
           "formulas": row["formulas"]}
    return hashlib.sha256(json.dumps(obj, sort_keys=True, ensure_ascii=False).encode()).hexdigest()


def read_state(path):
    if not path.exists():
        return {}
    with path.open(encoding="utf-8") as f:
        obj = json.load(f)
    if not isinstance(obj, dict):
        raise ValueError("Invalid state file")
    return obj


def save_state(path, state):
    tmp = path.with_name(path.name + ".tmp")
    with tmp.open("w", encoding="utf-8") as f:
        json.dump(state, f, indent=2, ensure_ascii=False)
        f.write("\n")
        f.flush()
        os.fsync(f.fileno())
    if os.name != "nt":
        os.chmod(tmp, 0o600)
    os.replace(tmp, path)


def csrf_from_jwt(token):
    try:
        parts = token.split(".")
        if len(parts) != 3:
            return None
        raw = base64.urlsafe_b64decode(parts[1] + "=" * (-len(parts[1]) % 4))
        payload = json.loads(raw.decode("utf-8"))
        return payload.get("csrfToken") or payload.get("csrf_token")
    except (ValueError, TypeError, UnicodeError, KeyError, binascii.Error):
        return None


def csrf_from_cookies(context, base_url):
    for c in context.cookies(base_url):
        name = c.get("name", "").lower()
        value = c.get("value", "")
        if name == "token":
            token = csrf_from_jwt(value)
            if isinstance(token, str) and token:
                return token
        if name in {"csrf_token", "csrftoken", "xsrf-token"} and value:
            return value
    return None


def request(page, method, path, payload=None, csrf=None):
    # Requests execute in the authenticated, same-origin browser context.
    # Only metadata is returned; no passwords, cookies or tokens are logged.
    return page.evaluate("""async a => {
      try {
        const headers = {'Accept': 'application/json'};
        if (a.method === 'POST') {
          headers['Content-Type'] = 'application/json';
          headers['X-CSRF-Token'] = a.csrf;
        }
        const r = await fetch(a.path, {method:a.method, credentials:'same-origin', headers,
          ...(a.method==='POST' ? {body:JSON.stringify(a.payload)} : {})});
        let j = {};
        if ((r.headers.get('content-type')||'').includes('application/json')) j = await r.json();
        const d = (j && typeof j.data==='object' && j.data!==null) ? j.data : {};
        return {status:r.status, codeS:j.codeS||null, role_id: d.unique_id||null,
                role_name:d.name||null};
      } catch(e) {return {status:null,codeS:null,role_id:null,role_name:null};}
    }""", {"method": method, "path": path, "payload": payload, "csrf": csrf})


def same_origin(base_url, current_url):
    a, b = urlparse(base_url), urlparse(current_url)
    return (a.scheme, a.hostname, a.port) == (b.scheme, b.hostname, b.port)


def plan(rows, state, only_user=None, limit=None):
    importable = [r for r in rows if r["people"]["mode"] == "import"]
    if only_user:
        importable = [r for r in importable if r["key"] == canon(only_user)]
        if not importable:
            raise ValueError("Selected username not in workbook or marked manual")
    blocked = [r for r in importable if r["reasons"]]
    if blocked:
        raise ValueError("Blocked input for " + ", ".join(r["people"]["username"] for r in blocked))
    pending, recorded = [], []
    for r in importable:
        previous = state.get(r["key"])
        if previous is None:
            pending.append(r)
        elif not isinstance(previous, dict) or previous.get("fingerprint") != fingerprint(r):
            raise ValueError("State does not match current workbook for " + r["key"])
        elif previous.get("stage") == "created_pending_manual_verification":
            recorded.append(r)
        else:
            raise ValueError("Incomplete/ambiguous previous attempt for " + r["key"] +
                             "; inspect People and state file. Never auto-retry")
    if limit is not None:
        if limit < 1:
            raise ValueError("--limit must be >= 1")
        pending = pending[:limit]
    return pending, recorded


def run_live(rows, args, config):
    from playwright.sync_api import sync_playwright
    base = config["base_url"].rstrip("/")
    if not base.startswith("https://") or not urlparse(base).hostname:
        raise ValueError("base_url must be absolute HTTPS URL")
    state_path = Path(args.state)
    state = read_state(state_path)
    pending, completed = plan(rows, state, args.user, args.limit)
    if not pending:
        print("No new users to create. Logged successes awaiting manual login/rights verification:", len(completed))
        return
    print("\n=== PLAN ===")
    for r in pending:
        print(f"  {r['people']['username']}: previous={r['legacy_username']} | {len(r['formulas'])} sites")
        for site, role in r["sites"]:
            print("     ", role, site)
    print("\nVERIFY IN PEOPLE that none of the target usernames already exist.")
    print("Only two POST requests per user: create role, then create local-only account.")
    print("This is an UNOFFICIAL API. No automatic retries, no deletion, no DB writes.")
    confirmation = f"IMPORT {len(pending)}"
    if input("To proceed, type exactly " + confirmation + ": ").strip() != confirmation:
        print("Cancelled; no server changes")
        return
    with sync_playwright() as pw:
        browser = pw.chromium.launch(channel=config.get("browser_channel", "chrome"), headless=False)
        context = browser.new_context(ignore_https_errors=bool(config.get("allow_insecure_tls", False)))
        page = context.new_page()
        seen = {"csrf": None}
        def on_response(response):
            if same_origin(base, response.url):
                token = response.headers.get("x-updated-csrf-token") or response.headers.get("x-csrf-token")
                if token:
                    seen["csrf"] = token
        page.on("response", on_response)
        try:
            page.goto(base, wait_until="domcontentloaded", timeout=45000)
            input("Log in as the LOCAL Owner in the opened browser, then press ENTER here ... ")
            if not same_origin(base, page.url):
                print("Unexpected browser origin. No changes")
                return
            get_result = request(page, "GET", READ_PATH)
            if get_result.get("status") != 200 or get_result.get("codeS") != "SUCCESS":
                print("Read API check failed. No changes")
                return
            csrf = csrf_from_cookies(context, base) or seen["csrf"]
            if not csrf:
                print("No CSRF token found. No changes")
                return
            last_role_second = 0
            for index, row in enumerate(pending, 1):
                key = row["key"]
                state = read_state(state_path)
                if key in state:
                    print("Stopped: prior state exists for", key)
                    return
                while int(time.time()) <= last_role_second:
                    time.sleep(0.2)
                last_role_second = int(time.time())
                role_name = f"role_private_{last_role_second}"
                role_payload = {"name": role_name, "is_private": True,
                                "permissions": [{"manifest_unique_key": "network.management",
                                                 "formulas": row["formulas"]}],
                                "permission_resources": []}
                state[key] = {"stage": "role_post_uncertain", "role_name": role_name,
                              "fingerprint": fingerprint(row), "when": time.strftime("%Y-%m-%d %H:%M:%S")}
                save_state(state_path, state)
                print(f"[{index}/{len(pending)}] Creating role for {row['people']['username']} ...", flush=True)
                csrf = seen["csrf"] or csrf_from_cookies(context, base) or csrf
                result = request(page, "POST", ROLE_PATH, role_payload, csrf)
                if (result.get("status") not in (200, 201) or result.get("codeS") != "SUCCESS"
                    or result.get("role_name") != role_name
                    or not UUID.fullmatch(result.get("role_id") or "")):
                    print("Role creation unconfirmed. HTTP:", result.get("status"),
                          "codeS:", result.get("codeS"), ". STOP; review People/state")
                    return
                role_id = result["role_id"]
                state[key].update({"stage": "role_created", "role_id": role_id})
                save_state(state_path, state)
                p = row["people"]
                user_payload = {
                    "first_name": p["first_name"], "last_name": p["last_name"],
                    "user_email": p["email"], "force_add_nfc": True,
                    "nfc_token": "", "group_ids": [], "pin_code": "",
                    "role_id": role_id, "username": p["username"],
                    "password": p["password"], "only_local_account": True,
                }
                state[key]["stage"] = "user_post_uncertain"
                save_state(state_path, state)
                print("  Creating local-only user ...", flush=True)
                csrf = seen["csrf"] or csrf_from_cookies(context, base) or csrf
                answer = request(page, "POST", USER_PATH, user_payload, csrf)
                del user_payload
                if answer.get("status") not in (200, 201) or answer.get("codeS") != "SUCCESS":
                    print("User creation unconfirmed. HTTP:", answer.get("status"),
                          "codeS:", answer.get("codeS"), ". STOP; do not retry blindly")
                    return
                state[key]["stage"] = "created_pending_manual_verification"
                save_state(state_path, state)
                print("  API confirmed. Verify login and visible site permissions manually")
            print("\nAPI accepted all new users. This does NOT prove each account's effective permissions.")
        finally:
            input("Press ENTER to close the test browser ... ")
            context.close()
            browser.close()


def main():
    parser = argparse.ArgumentParser(description="UNOFFICIAL UniFi OS multi-site local admin migration")
    parser.add_argument("--config", default="config.json")
    parser.add_argument("--legacy", default="legacy_privileges.jsonl")
    parser.add_argument("--people", default="administrators.xlsx")
    parser.add_argument("--target-sites", default="target_sites.jsonl")
    parser.add_argument("--state", default="import_state.json")
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument("--preview", action="store_true")
    group.add_argument("--user", help="Import just one named username")
    group.add_argument("--all", action="store_true", help="Import all uncreated import-mode rows")
    parser.add_argument("--limit", type=int, help="Optional cap for a small trial batch")
    args = parser.parse_args()
    config = json.loads(Path(args.config).read_text(encoding="utf-8"))
    people = read_people(Path(args.people))
    legacy, counts = read_legacy(Path(args.legacy))
    target = read_target_sites(Path(args.target_sites))
    manual = {canon(x) for x in config.get("manual_admin_usernames", [])}
    rows = prepare(people, legacy, target, manual)
    print(f"Legacy admins: {len(legacy)} | privilege records: {sum(counts.values())} "
          f"| target sites: {len(target)} | workbook users: {len(people)}")
    print("Legacy role counts:", dict(counts))
    if args.preview:
        print("\n=== OFFLINE PREVIEW: no server connection, no secrets printed ===")
        for row in rows:
            p = row["people"]
            state = read_state(Path(args.state)).get(row["key"])
            status = ("MANUAL" if p["mode"] == "manual" else
                      "BLOCKED" if row["reasons"] else
                      "RECORDED" if state and state.get("stage") == "created_pending_manual_verification" else
                      "HOLD_REVIEW" if state else "READY")
            print(f"{p['username']:<24} legacy={row['legacy_username']:<22} "
                  f"sites={len(row['formulas']):>3} {status}")
            for problem in row["reasons"]:
                print("    ERROR:", problem)
            for warning in row["warnings"]:
                print("    CHECK:", warning)
        print("This preview does not verify accounts already present in People.")
        return
    if any(r["reasons"] for r in rows):
        raise ValueError("Resolve every blocked workbook row before importing")
    if args.limit is not None and args.user:
        raise ValueError("Do not combine --limit with --user")
    run_live(rows, args, config)


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError, KeyError, json.JSONDecodeError) as exc:
        print("STOPPED:", str(exc)[:250])
        raise SystemExit(1)
